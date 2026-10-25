"""Executable checks for the declared DOCX and PDF document boundaries."""

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.document_parser import extract_text_from_docx, extract_text_from_pdf
from tests.conftest import get_auth_headers


def test_docx_parser_extracts_paragraphs_and_table_cells():
    stream = BytesIO()
    document = Document()
    document.add_paragraph("No pleural effusion.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Heart rate"
    table.cell(0, 1).text = "82 bpm"
    document.save(stream)

    extracted = extract_text_from_docx(stream.getvalue())

    assert "No pleural effusion." in extracted
    assert "Heart rate | 82 bpm" in extracted


def _blank_pdf() -> bytes:
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(stream)
    return stream.getvalue()


def test_blank_scanned_page_is_rejected_after_ocr_finds_no_text():
    with pytest.raises(ValueError, match="no readable text"):
        extract_text_from_pdf(_blank_pdf())


def test_without_ocr_an_image_only_pdf_is_rejected_explicitly(monkeypatch):
    from app.services import ocr

    monkeypatch.setattr(ocr, "ocr_available", lambda: False)
    with pytest.raises(ValueError, match="OCR is not configured"):
        extract_text_from_pdf(_blank_pdf())


def test_docx_upload_runs_real_parser_and_persists_measurement(client, test_users):
    headers = get_auth_headers(test_users["clinician_a"])
    case_response = client.post(
        "/api/v1/cases",
        json={
            "pseudonymous_id": "CASE-DOCX-PARSER",
            "examinations": [{
                "exam_date": "2025-03-10T12:00:00Z",
                "modality": "REPORT",
                "clinical_indication": "Synthetic parser fixture",
            }],
        },
        headers=headers,
    )
    assert case_response.status_code == 201
    case_record = case_response.json()

    stream = BytesIO()
    document = Document()
    document.add_paragraph("Heart rate: 82 bpm. No pleural effusion.")
    document.save(stream)
    upload_response = client.post(
        f"/api/v1/cases/{case_record['id']}/uploads",
        files={
            "file": (
                "deidentified-report.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={
            "examination_id": str(case_record["examinations"][0]["id"]),
            "is_associated": "true",
            "deidentified": "true",
        },
        headers=headers,
    )
    assert upload_response.status_code == 201

    job_response = client.post(
        f"/api/v1/uploads/{upload_response.json()['id']}/analysis-jobs",
        headers=headers,
    )
    assert job_response.status_code == 201, job_response.text
    assert job_response.json()["status"] == "COMPLETED"

    measurements = client.get(
        f"/api/v1/cases/{case_record['id']}/measurements",
        headers=headers,
    ).json()
    assert len(measurements) == 1
    assert measurements[0]["name"] == "heart_rate"
    assert measurements[0]["value"] == 82.0
    assert measurements[0]["unit"] == "bpm"
    assert measurements[0]["source_span"] == "Heart rate: 82 bpm"
